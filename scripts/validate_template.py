"""Deterministic gate for generated spreadsheet templates.

A template is two artifacts that must agree:
  * metadata JSON in <templatesTarget.contentDir>/<slug>.json
  * the workbook itself in <templatesTarget.filesDir>/<slug>.xlsx

Metadata checks are stdlib. Workbook checks need openpyxl; point
BLOG_PIPELINE_PYTHON at an interpreter that has it (a venv works) or run this
script with one. Without openpyxl the workbook checks fail loudly rather than
silently passing an unopenable file.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_SCRIPTS_DIR = str(Path(__file__).resolve().parent)
if _SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, _SCRIPTS_DIR)

from config import load_config, ConfigError

REQUIRED_FIELDS = ("slug", "title", "category", "shortDescription", "about", "link", "order")

# Same prose-cleanliness rule as blogs: human copy, no AI-tell punctuation.
BANNED_SEQUENCES = {
    "—": "em dash", "–": "en dash",
    "’": "curly apostrophe", "‘": "curly quote",
    "“": "curly quote", "”": "curly quote",
}


def validate_metadata(path: Path, cfg: dict, existing_slugs: list) -> list:
    errors: list = []
    try:
        data = json.loads(path.read_text())
    except (json.JSONDecodeError, OSError) as exc:
        return [f"metadata unreadable: {exc}"]

    for field in REQUIRED_FIELDS:
        if field not in data or data[field] in (None, ""):
            errors.append(f"missing field: {field}")
    if data.get("slug") != path.stem:
        errors.append(f"slug {data.get('slug')!r} != filename {path.stem!r}")
    if data.get("slug") in existing_slugs:
        errors.append(f"duplicate slug: {data.get('slug')}")

    tcfg = cfg.get("templatesTarget", {})
    cats = tcfg.get("categories")
    if cats and data.get("category") not in cats:
        errors.append(f"category {data.get('category')!r} not in templatesTarget.categories")

    link = str(data.get("link", ""))
    if link and not (link.startswith("/") or link.startswith("https://")):
        errors.append(f"link must be a site-relative path or https URL, got {link!r}")

    for field in ("shortDescription", "about"):
        text = str(data.get(field, ""))
        for seq, name in BANNED_SEQUENCES.items():
            if seq in text:
                errors.append(f"{field}: contains {name} x{text.count(seq)}")
    return errors


def validate_workbook(path: Path) -> list:
    if not path.exists():
        return [f"workbook missing: {path}"]
    try:
        import openpyxl
    except ImportError:
        return [
            "openpyxl not importable; run with an interpreter that has it "
            "(e.g. ~/.blog-pipeline-venv/bin/python)"
        ]
    errors: list = []
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as exc:
        return [f"workbook does not open: {exc}"]

    names = [s.lower() for s in wb.sheetnames]
    if len(wb.sheetnames) < 2:
        errors.append("workbook needs at least 2 sheets (Instructions + the model)")
    if not any("instruction" in n or "read me" in n or "readme" in n for n in names):
        errors.append("no Instructions sheet (first sheet should tell the user how to use it)")

    has_formula = False
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    has_formula = True
                    break
            if has_formula:
                break
        if has_formula:
            break
    if not has_formula:
        errors.append("no formulas found; a template should compute, not just hold labels")
    return errors


def validate_template(repo_path: str, slug: str) -> list:
    cfg = load_config(repo_path)
    tcfg = cfg.get("templatesTarget")
    if not tcfg:
        return ["config has no templatesTarget section; add it (contentDir, filesDir, categories)"]
    repo = Path(repo_path)
    content_dir = repo / tcfg["contentDir"]
    files_dir = repo / tcfg["filesDir"]
    existing = [p.stem for p in content_dir.glob("*.json") if p.stem != slug]
    errors = validate_metadata(content_dir / f"{slug}.json", cfg, existing)
    errors += validate_workbook(files_dir / f"{slug}.xlsx")
    return errors


def main(argv):
    if len(argv) < 3 or argv[1] in ("-h", "--help"):
        print("usage: validate_template.py <target-repo-path> <slug> [more slugs...]")
        return 2
    try:
        failed = False
        for slug in argv[2:]:
            errors = validate_template(argv[1], slug)
            if errors:
                failed = True
                print(f"FAIL {slug}")
                for e in errors:
                    print(f"  - {e}")
            else:
                print(f"OK   {slug}")
        return 1 if failed else 0
    except ConfigError as exc:
        print(f"config invalid:\n{exc}")
        return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
